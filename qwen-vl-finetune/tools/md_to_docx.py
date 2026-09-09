#!/usr/bin/env python3
"""Render a Markdown document as .docx (upload to Google Drive and open with
Google Docs: headings, lists, tables, bold and embedded images survive).

Supported Markdown subset: #/##/### headings, paragraphs (wrapped lines are
joined), *bullet* and 1. numbered lists with two-space nesting, pipe tables
(first row = header), ![alt](image.png) images (paths relative to the .md),
and **bold** / *italic* / `code` inline markup.

Usage:
    python tools/md_to_docx.py docs/proposal/MLPerf_Training_proposal_Qwen3-VL.md
    (writes the .docx next to the .md, or pass --out)
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
LIST_ITEM = re.compile(r"^(\s*)([*+-]|\d+\.)\s+(.*)$")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*))")


def add_inline(paragraph, text):
    """Add text to a paragraph, honoring **bold**, *italic* and `code`."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def new_numbering_instance(doc, style_name="List Number"):
    """Create a fresh numbering instance (restarting at 1) that reuses the
    style's list definition. Word shares one instance per style, so without
    this every numbered list in the document continues the previous count."""
    try:
        style_num_id = doc.styles[style_name].element.pPr.numPr.numId.val
    except AttributeError:
        return None
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    nums = numbering.findall(qn("w:num"))
    abstract_id = next(
        n.find(qn("w:abstractNumId")).get(qn("w:val"))
        for n in nums
        if n.get(qn("w:numId")) == str(style_num_id)
    )
    new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def set_numbering(paragraph, num_id, level=0):
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    paragraph._p.get_or_add_pPr().append(num_pr)


def add_table(doc, rows, total_width=Inches(6.5)):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    table.autofit = False
    font_size = Pt(8) if width > 6 else Pt(9)

    # Column widths: each column gets at least the rendered width of its longest
    # word (so headers don't break mid-word); remaining width is shared in
    # proportion to the longest cell text. Falls back to proportional scaling
    # when even the minimums don't fit.
    def longest(j, whole_cell):
        texts = [re.sub(r"[*`]", "", r[j]) for r in cells if j < len(r)]
        if not whole_cell:
            texts = [w for t in texts for w in t.split()]
        return max((len(t) for t in texts), default=1)

    char_width = 0.08 if font_size == Pt(8) else 0.09  # inches per bold character (approx.)
    total_in = total_width / 914400  # EMU -> inches
    minimums = [longest(j, False) * char_width + 0.16 for j in range(width)]
    if sum(minimums) <= total_in:
        full = [longest(j, True) for j in range(width)]
        extra = total_in - sum(minimums)
        widths_in = [m + extra * f / sum(full) for m, f in zip(minimums, full)]
    else:
        widths_in = [total_in * m / sum(minimums) for m in minimums]
    col_widths = [Inches(w) for w in widths_in]
    for j, col_width in enumerate(col_widths):
        table.columns[j].width = col_width  # grid widths: honored by LibreOffice/Docs

    for i, row in enumerate(cells):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))  # keep a row on one page
        for j in range(width):
            cell = table.cell(i, j)
            cell.width = col_widths[j]
            cell.text = ""
            para = cell.paragraphs[0]
            add_inline(para, row[j] if j < len(row) else "")
            for run in para.runs:
                run.font.size = font_size
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def flush_paragraph(doc, lines, base_dir, state):
    """Emit the accumulated lines as one list item, image or paragraph.
    `state["num_id"]` tracks the numbering instance of the numbered list in
    progress; it is reset by any non-list content so the next list restarts."""
    if not lines:
        return
    first = lines[0]
    match = LIST_ITEM.match(first)
    text = " ".join(s.strip() for s in lines)
    if match:
        indent, marker, _ = match.groups()
        level = min(len(indent) // 2, 2)
        numbered = marker[0].isdigit()
        style = ("List Number" if numbered else "List Bullet") + (f" {level + 1}" if level else "")
        text = " ".join([match.group(3)] + [s.strip() for s in lines[1:]])
        paragraph = doc.add_paragraph(style=style)
        add_inline(paragraph, text)
        if numbered and level == 0:
            if state.get("num_id") is None:
                state["num_id"] = new_numbering_instance(doc)
            if state["num_id"] is not None:
                set_numbering(paragraph, state["num_id"])
        elif not numbered and level == 0:
            state["num_id"] = None
        return
    state["num_id"] = None
    if text.startswith("![") and text.endswith(")"):
        path = text[text.rfind("](") + 2 : -1]
        doc.add_picture(str((base_dir / path).resolve()), width=Inches(6.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        return
    add_inline(doc.add_paragraph(), text)


def convert(md_path, out_path):
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.75)
    text_width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    base_dir = md_path.parent
    lines = md_path.read_text().splitlines()

    para, table, state = [], [], {"num_id": None}
    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("|"):
            flush_paragraph(doc, para, base_dir, state)
            para = []
            if not TABLE_SEP.match(stripped):
                table.append(stripped)
            continue
        if table:
            add_table(doc, table, text_width)
            table = []
            state["num_id"] = None
        if not stripped:
            flush_paragraph(doc, para, base_dir, state)
            para = []
        elif stripped.startswith("#"):
            flush_paragraph(doc, para, base_dir, state)
            para = []
            state["num_id"] = None
            level = len(stripped) - len(stripped.lstrip("#"))
            add_inline(doc.add_heading(level=min(level, 3)), stripped.lstrip("#").strip())
        elif LIST_ITEM.match(line) and para:
            flush_paragraph(doc, para, base_dir, state)  # new list item starts a new paragraph
            para = [line]
        else:
            para.append(line)
    if table:
        add_table(doc, table, text_width)
    flush_paragraph(doc, para, base_dir, state)
    doc.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", help="Input .md file")
    parser.add_argument("--out", help="Output .docx path (default: alongside the input)")
    args = parser.parse_args()
    src = Path(args.markdown)
    dst = Path(args.out) if args.out else src.with_suffix(".docx")
    convert(src, dst)
    print(f"wrote {dst}")


if __name__ == "__main__":
    main()
